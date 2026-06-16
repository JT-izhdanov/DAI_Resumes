"""Extract text from résumé files and resolve résumé pools.

A *pool* is a source of résumés to assess. It can be:
  - a role name        → that role's roles/<role>/resumes/ folder
  - a directory path   → every supported file in it
  - a glob path        → matching files

Decoupling pools from the role being scored is what lets you assess, say, the
Data Engineer applicant pool against the Senior Data Engineer rubric.
"""

from __future__ import annotations

import glob
from pathlib import Path
from typing import List, NamedTuple

from . import config
from . import criteria as criteria_mod


class ResumeRef(NamedTuple):
    """A résumé file plus the pool it came from (for provenance)."""

    path: Path
    source: str  # pool label, e.g. the role the résumé was submitted for


def extract_text(path: Path) -> str:
    """Return the plain text of a résumé file (.pdf or .docx)."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    raise ValueError(
        f"Unsupported file type '{suffix}' for {path.name}. "
        f"Supported: {', '.join(sorted(config.SUPPORTED_EXTENSIONS))}"
    )


def _extract_pdf(path: Path) -> str:
    import pdfplumber

    parts: List[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
    return "\n\n".join(parts).strip()


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:  # résumés often use tables for layout
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _supported_files(paths) -> List[Path]:
    return sorted(
        p for p in paths if p.is_file() and p.suffix.lower() in config.SUPPORTED_EXTENSIONS
    )


def list_resumes(role: str) -> List[Path]:
    """All parseable résumés in a role's own pool."""
    folder = config.role_resume_dir(role)
    if not folder.exists():
        return []
    return _supported_files(folder.iterdir())


def gather(pools: List[str]) -> List[ResumeRef]:
    """Resolve and de-duplicate résumés across one or more pool specs."""
    refs: List[ResumeRef] = []
    seen: set[Path] = set()
    for spec in pools:
        label, files = _resolve_one(spec)
        for f in files:
            resolved = f.resolve()
            if resolved in seen:
                continue
            seen.add(resolved)
            refs.append(ResumeRef(path=f, source=label))
    return refs


def _resolve_one(spec: str) -> tuple[str, List[Path]]:
    # 1. A known role → its résumé pool.
    if criteria_mod.is_role(spec):
        return spec, list_resumes(spec)
    # 2. A directory on disk.
    p = Path(spec)
    if p.is_dir():
        return p.name, _supported_files(p.iterdir())
    # 3. A glob.
    matches = [Path(m) for m in glob.glob(spec)]
    if matches:
        return Path(spec).parent.name or "files", _supported_files(matches)
    raise FileNotFoundError(
        f"Pool '{spec}' is not a known role, a directory, or a matching glob."
    )
